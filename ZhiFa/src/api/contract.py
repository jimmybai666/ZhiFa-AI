"""
合同审查API模块
"""
from flask import request, jsonify, send_file, after_this_request, Response
import traceback
import os
import uuid
import json
import shutil
import threading
from werkzeug.utils import secure_filename

from ..core.cache_manager import increment_stat
from ..vectorstore.utils import get_vectorstore
from ..utils.file_converter import convert_to_markdown
from ..utils.export_utils import generate_pdf_report, generate_revised_docx, generate_redline_html, create_export_zip
from ..config.config import Config
from ..chain.contract_review_prompt import CONTRACT_STRUCTURE_PROMPT, CONTRACT_CLAUSE_ANALYSIS_PROMPT, CONTRACT_SUMMARY_PROMPT
from ..core.model_factory import get_model
from ..utils.redline_diff import compute_redline_diff

# 活跃的审查会话取消信号 {session_id: threading.Event}
_active_reviews = {}


def register_contract_routes(app, get_or_refresh_contract_chain):
    """注册合同审查相关路由"""

    @app.route('/api/contract-review/cancel/<session_id>', methods=['POST'])
    def cancel_contract_review(session_id):
        """取消正在进行的合同审查"""
        cancel_event = _active_reviews.get(session_id)
        if cancel_event:
            cancel_event.set()
            return jsonify({"status": "cancelled"})
        return jsonify({"status": "not_found"}), 404
    @app.route('/api/contract-review', methods=['POST'])
    def contract_review():
        """
        合同审查API
        """
        try:
            contract_text = ""
            session_id = str(uuid.uuid4())
            original_file_path = None
            original_filename = None

            # 1. 尝试处理文件上传
            if 'file' in request.files:
                file = request.files['file']
                if file and file.filename:
                    filename = secure_filename(file.filename)
                    original_filename = filename
                    
                    # 确保临时目录存在
                    temp_dir = os.path.join(os.getcwd(), 'storage', 'temp_uploads', session_id)
                    os.makedirs(temp_dir, exist_ok=True)
                    
                    file_path = os.path.join(temp_dir, filename)
                    file.save(file_path)
                    original_file_path = file_path
                    
                    try:
                        # 转换文件为 Markdown/文本
                        # TODO: convert_to_markdown (PaddleOCR) is a blocking call and cannot be easily interrupted.
                        # To support cancellation, we would need to run this in a separate process/thread
                        # and implement a mechanism to kill it.
                        contract_text = convert_to_markdown(file_path)
                    except Exception as e:
                        return jsonify({"error": f"文件解析失败: {str(e)}"}), 400
                    # 注意：这里不再立即删除文件，而是保留供导出使用

            # 2. 如果没有文件或文件解析为空，尝试从 JSON 获取
            if not contract_text:
                # silent=True 防止 Content-Type 不是 application/json 时报错
                data = request.get_json(silent=True)
                if data:
                    contract_text = data.get('contract_text', '')
                    # 如果是纯文本输入，我们创建一个临时的 docx 作为"原始文件"
                    if contract_text:
                        temp_dir = os.path.join(os.getcwd(), 'storage', 'temp_uploads', session_id)
                        os.makedirs(temp_dir, exist_ok=True)
                        original_filename = "original_contract.docx"
                        original_file_path = os.path.join(temp_dir, original_filename)
                        
                        # 使用 python-docx 创建简单的文档
                        from docx import Document
                        from docx.oxml.ns import qn
                        doc = Document()
                        
                        # 设置默认中文字体
                        style = doc.styles['Normal']
                        style.font.name = 'Microsoft YaHei'
                        style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
                        
                        for line in contract_text.split('\n'):
                            doc.add_paragraph(line)
                        doc.save(original_file_path)
            
            if not contract_text:
                return jsonify({"error": "请提供合同文本或上传合同文件"}), 400
            
            # 获取最新的 chain（支持热更新）
            contract_chain = get_or_refresh_contract_chain()
            
            if contract_chain is None:
                return jsonify({"error": "合同审查服务未初始化"}), 500
            
            # 调用合同审查链
            result = contract_chain.invoke({"contract_text": contract_text})
            
            # 更新使用统计
            increment_stat('contract_count')
            
            return jsonify({
                "clauses_analysis": result.get("clauses_analysis", []),
                "summary_report": result.get("summary_report", ""),
                "total_clauses": result.get("total_clauses", 0),
                "session_id": session_id,
                "original_filename": original_filename
            })
            
        except Exception as e:
            traceback.print_exc()
            return jsonify({"error": str(e)}), 500

    @app.route('/api/contract-review/stream', methods=['POST'])
    def contract_review_stream():
        """
        合同审查SSE流式接口 - 实时推送进度
        """
        try:
            contract_text = ""
            session_id = str(uuid.uuid4())
            original_file_path = None
            original_filename = None

            if 'file' in request.files:
                file = request.files['file']
                if file and file.filename:
                    filename = secure_filename(file.filename)
                    original_filename = filename
                    temp_dir = os.path.join(os.getcwd(), 'storage', 'temp_uploads', session_id)
                    os.makedirs(temp_dir, exist_ok=True)
                    file_path = os.path.join(temp_dir, filename)
                    file.save(file_path)
                    original_file_path = file_path

            if not contract_text:
                if original_file_path:
                    pass
                else:
                    data = request.get_json(silent=True)
                    if data:
                        contract_text = data.get('contract_text', '')
                        if contract_text:
                            temp_dir = os.path.join(os.getcwd(), 'storage', 'temp_uploads', session_id)
                            os.makedirs(temp_dir, exist_ok=True)
                            original_filename = "original_contract.docx"
                            original_file_path = os.path.join(temp_dir, original_filename)
                            from docx import Document as DocxDocument
                            from docx.oxml.ns import qn
                            doc = DocxDocument()
                            style = doc.styles['Normal']
                            style.font.name = 'Microsoft YaHei'
                            style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
                            for line in contract_text.split('\n'):
                                doc.add_paragraph(line)
                            doc.save(original_file_path)

            if not contract_text and not original_file_path:
                return jsonify({"error": "请提供合同文本或上传合同文件"}), 400

            req_contract_text = contract_text
            req_original_file_path = original_file_path
            req_session_id = session_id
            req_original_filename = original_filename

            # 注册取消信号
            cancel_event = threading.Event()
            _active_reviews[session_id] = cancel_event

            def generate():
                import time as _time
                from concurrent.futures import ThreadPoolExecutor, as_completed

                nonlocal req_contract_text, req_original_file_path

                def send_event(data):
                    return f"data: {json.dumps(data, ensure_ascii=False)}\n\n"

                def is_cancelled():
                    return cancel_event.is_set()

                try:
                    # 首先发送 session_id 供前端用于取消
                    yield send_event({"type": "session", "session_id": req_session_id})

                    # Step 0: 解析文档
                    yield send_event({"type": "progress", "step": 0})

                    if req_original_file_path and not req_contract_text:
                        try:
                            req_contract_text = convert_to_markdown(req_original_file_path)
                        except Exception as e:
                            yield send_event({"type": "error", "message": f"文件解析失败: {str(e)}"})
                            return

                    if not req_contract_text:
                        yield send_event({"type": "error", "message": "合同文本为空"})
                        return

                    if is_cancelled():
                        yield send_event({"type": "cancelled"})
                        return

                    # Step 1: 拆分条款 (LLM结构化)
                    yield send_event({"type": "progress", "step": 1})

                    config = Config()
                    law_vs = get_vectorstore(config.LAW_VS_COLLECTION_NAME)
                    vs_retriever = law_vs.as_retriever(search_kwargs={"k": 5})
                    import re as _re

                    model = get_model(streaming=False)
                    prompt = CONTRACT_STRUCTURE_PROMPT.format(contract_text=req_contract_text)
                    try:
                        response = model.invoke(prompt)
                        content = response.content if hasattr(response, 'content') else str(response)
                        content = content.strip()
                        if content.startswith("```json"):
                            content = content[7:]
                        elif content.startswith("```"):
                            content = content[3:]
                        if content.endswith("```"):
                            content = content[:-3]
                        content = content.strip()
                        clauses_data = json.loads(content)
                        for item in clauses_data:
                            if "content" in item and isinstance(item["content"], str):
                                item["content"] = _re.sub(r'(?m)^#+\s*', '', item["content"])
                    except Exception:
                        clauses = _re.split(r'\n\n+|\n', req_contract_text)
                        clauses_data = [
                            {"id": str(i), "type": "paragraph", "theme": "未分类", "section": "自动分割", "content": c.strip()}
                            for i, c in enumerate(clauses, 1) if c.strip()
                        ]

                    # Step 2: 风险检测 (并发分析条款)
                    if is_cancelled():
                        yield send_event({"type": "cancelled"})
                        return
                    yield send_event({"type": "progress", "step": 2})

                    def combine_law_docs(docs):
                        return "\n\n".join([d.page_content for d in docs])

                    def analyze_clause(item):
                        if is_cancelled():
                            return None
                        clause_text = item.get("content", "") or "（无内容）"
                        clause_id = item.get("id", "unknown")
                        clause_theme = item.get("theme", "未分类")
                        section_path = item.get("section", "")
                        law_docs = vs_retriever.invoke(clause_text)
                        law_context = combine_law_docs(law_docs)
                        evidence_list = [{"content": d.page_content, "source": d.metadata.get("source", "Unknown"), "page": d.metadata.get("page", 0)} for d in law_docs]
                        if not law_context.strip():
                            law_context = "未检索到直接相关的法律条文"
                        structured_clause_text = f"【章节：{section_path}】【主题：{clause_theme}】\n内容：{clause_text}"
                        prompt_value = CONTRACT_CLAUSE_ANALYSIS_PROMPT.format_prompt(clause=structured_clause_text, law_context=law_context)
                        if is_cancelled():
                            return None
                        m = get_model(streaming=False)
                        analysis_result = m.invoke(prompt_value.to_string())
                        result_content = analysis_result.content if hasattr(analysis_result, 'content') else str(analysis_result)
                        try:
                            clean_json = result_content.strip()
                            if clean_json.startswith("```json"):
                                clean_json = clean_json[7:]
                            elif clean_json.startswith("```"):
                                clean_json = clean_json[3:]
                            if clean_json.endswith("```"):
                                clean_json = clean_json[:-3]
                            analysis_json = json.loads(clean_json.strip())
                            risk_level = analysis_json.get("risk_level", "未知")
                            issue_description = analysis_json.get("issue_description", "")
                            suggestion = analysis_json.get("suggestion", "")
                            revised_text = analysis_json.get("revised_text", clause_text)
                            diff_html = compute_redline_diff(clause_text, revised_text)
                            analysis_text = f"风险等级：{risk_level}\n问题：{issue_description}\n建议：{suggestion}"
                        except json.JSONDecodeError:
                            analysis_text = result_content
                            risk_level = "解析失败"
                            diff_html = clause_text
                            evidence_list = []
                            revised_text = clause_text
                        return {
                            "clause_number": clause_id, "clause_text": clause_text,
                            "type": item.get("type", "条款"), "theme": clause_theme,
                            "section": section_path, "analysis": analysis_text,
                            "risk_level": risk_level, "diff_html": diff_html,
                            "revised_text": revised_text, "evidence": evidence_list,
                            "related_laws": law_context
                        }

                    total = len(clauses_data)
                    clauses_analysis = [None] * total
                    with ThreadPoolExecutor(max_workers=4) as executor:
                        futures = [executor.submit(analyze_clause, item) for item in clauses_data]
                        for future in as_completed(futures):
                            if is_cancelled():
                                executor.shutdown(wait=False, cancel_futures=True)
                                break
                            idx_result = future.result()
                            if idx_result is None:
                                continue
                            for i, item in enumerate(clauses_data):
                                if item.get("id") == idx_result["clause_number"]:
                                    clauses_analysis[i] = idx_result
                                    break

                    if is_cancelled():
                        yield send_event({"type": "cancelled"})
                        return

                    # Step 3: 生成报告
                    yield send_event({"type": "progress", "step": 3})

                    clauses_text = "\n\n".join([
                        f"【条款 {item['clause_number']} - {item['theme']}】\n原文：{item['clause_text']}\n分析：{item['analysis']}"
                        for item in clauses_analysis if item
                    ])
                    summary_prompt = CONTRACT_SUMMARY_PROMPT.format_prompt(clauses_analysis=clauses_text)
                    summary_result = model.invoke(summary_prompt.to_string())
                    summary_text = summary_result.content if hasattr(summary_result, 'content') else str(summary_result)

                    increment_stat('contract_count')

                    yield send_event({
                        "type": "result",
                        "data": {
                            "clauses_analysis": [c for c in clauses_analysis if c],
                            "summary_report": summary_text,
                            "total_clauses": total,
                            "session_id": req_session_id,
                            "original_filename": req_original_filename
                        }
                    })
                finally:
                    _active_reviews.pop(req_session_id, None)

            return Response(generate(), mimetype='text/event-stream', headers={
                'Cache-Control': 'no-cache',
                'X-Accel-Buffering': 'no'
            })

        except Exception as e:
            traceback.print_exc()
            return jsonify({"error": str(e)}), 500

    @app.route('/api/export-contract-review', methods=['POST'])
    def export_contract_review():
        """
        导出合同审查结果为 ZIP
        """
        try:
            data = request.get_json()
            session_id = data.get('session_id')
            clauses_analysis = data.get('clauses_analysis', [])
            summary_report = data.get('summary_report', '')
            
            if not session_id:
                return jsonify({"error": "Missing session_id"}), 400
                
            # 临时目录
            base_temp_dir = os.path.join(os.getcwd(), 'storage', 'temp_uploads', session_id)
            export_dir = os.path.join(base_temp_dir, 'export')
            os.makedirs(export_dir, exist_ok=True)
            
            # 1. 生成 PDF 报告
            pdf_path = os.path.join(export_dir, '1. 审查报告.pdf')
            # 组合报告内容
            full_report_md = f"# 合同审查报告\n\n{summary_report}\n\n## 详细条款分析\n\n"
            for clause in clauses_analysis:
                risk = clause.get('risk_level', '无风险')
                if "高" in risk or "中" in risk:
                    full_report_md += f"### {clause.get('section', '条款')} ({risk})\n"
                    full_report_md += f"**原文**：{clause.get('clause_text')}\n\n"
                    full_report_md += f"**分析**：{clause.get('analysis')}\n\n"
            
            generate_pdf_report(full_report_md, pdf_path)
            
            # 2. 生成修订后合同 DOCX
            docx_path = os.path.join(export_dir, '2. 修订后合同.docx')
            generate_revised_docx(clauses_analysis, docx_path)
            
            # 3. 生成修订对比 HTML
            html_path = os.path.join(export_dir, '3. 修订对比.html')
            generate_redline_html(clauses_analysis, html_path)
            
            # 4. 查找原始文件
            original_file_path = None
            # 遍历 session 目录找原始文件 (排除 export 目录)
            for f in os.listdir(base_temp_dir):
                if f != 'export':
                    original_file_path = os.path.join(base_temp_dir, f)
                    break
            
            # 5. 打包 ZIP
            zip_filename = f"contract_review_{session_id}.zip"
            zip_path = os.path.join(base_temp_dir, zip_filename)
            create_export_zip(export_dir, zip_path, original_file_path)
            
            # 发送文件
            @after_this_request
            def remove_file(response):
                try:
                    # 清理整个 session 目录
                    shutil.rmtree(base_temp_dir)
                except Exception as error:
                    print(f"Error removing temp dir: {error}")
                return response
                
            return send_file(zip_path, as_attachment=True, download_name=zip_filename)

        except Exception as e:
            traceback.print_exc()
            return jsonify({"error": str(e)}), 500

