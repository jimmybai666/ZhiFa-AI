import os
import zipfile
import json
import shutil
import pathlib
import tempfile
import base64
import re
import html
from docx import Document
from xhtml2pdf import pisa
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
import markdown
from datetime import datetime

def generate_pdf_report(report_content, output_path):
    """
    Generates a PDF report from markdown content using ReportLab directly.
    This avoids xhtml2pdf issues with temp files and paths on Windows.
    """
    try:
        # 1. Register Chinese Font
        font_name = 'SimHei'
        font_path = r"C:\Windows\Fonts\simhei.ttf"
        
        # Fallback fonts
        if not os.path.exists(font_path):
            candidates = [
                r"C:\Windows\Fonts\msyh.ttf",
                os.path.join(os.getcwd(), 'storage', 'fonts', 'simhei.ttf')
            ]
            for path in candidates:
                if os.path.exists(path):
                    font_path = path
                    break
        
        try:
            pdfmetrics.registerFont(TTFont(font_name, font_path))
        except Exception as e:
            print(f"Font registration warning: {e}")
            # If registration fails, we might still be able to proceed if it was already registered
            pass

        # 2. Setup Document
        doc = SimpleDocTemplate(
            output_path, 
            pagesize=A4,
            rightMargin=72, leftMargin=72,
            topMargin=72, bottomMargin=18
        )
        
        # 3. Define Styles
        styles = getSampleStyleSheet()
        
        # Normal Text Style
        style_normal = ParagraphStyle(
            'ChineseNormal',
            parent=styles['Normal'],
            fontName=font_name,
            fontSize=10.5,
            leading=18,
            spaceAfter=12,
            firstLineIndent=21  # 2 chars indent
        )
        
        # Heading Styles
        style_h1 = ParagraphStyle(
            'ChineseH1',
            parent=styles['Heading1'],
            fontName=font_name,
            fontSize=16,
            leading=24,
            spaceAfter=16,
            spaceBefore=12,
            textColor=colors.HexColor('#2c3e50')
        )
        
        style_h2 = ParagraphStyle(
            'ChineseH2',
            parent=styles['Heading2'],
            fontName=font_name,
            fontSize=14,
            leading=20,
            spaceAfter=12,
            spaceBefore=10,
            textColor=colors.HexColor('#34495e')
        )
        
        style_h3 = ParagraphStyle(
            'ChineseH3',
            parent=styles['Heading3'],
            fontName=font_name,
            fontSize=12,
            leading=18,
            spaceAfter=10,
            spaceBefore=8,
            textColor=colors.HexColor('#455a64')
        )

        # 4. Parse Content
        story = []
        
        # Remove manual title addition since it's usually in the markdown content
        # story.append(Paragraph("合同审查报告", style_h1))
        # story.append(Spacer(1, 12))
        
        lines = report_content.split('\n')
        for line in lines:
            line = line.strip()
            if not line:
                continue
                
            # Handle Markdown Headers
            if line.startswith('# '):
                story.append(Paragraph(line[2:], style_h1))
            elif line.startswith('## '):
                story.append(Paragraph(line[3:], style_h2))
            elif line.startswith('### '):
                story.append(Paragraph(line[4:], style_h3))
            elif line.startswith('#### '):
                story.append(Paragraph(line[5:], style_h3))
            elif line.startswith('##### '):
                story.append(Paragraph(line[6:], style_h3))
            # Handle Lists
            elif line.startswith('- ') or line.startswith('* '):
                # Replace bold syntax **text** with <b>text</b>
                text = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', line[2:])
                # Use a font-safe bullet point. SimHei might not have U+2022.
                # U+25CF (Black Circle) or U+00B7 (Middle Dot) are usually safe.
                # We use U+25CF for a solid bullet look.
                story.append(Paragraph(f"● {text}", style_normal))
            # Normal Text
            else:
                # Replace bold syntax
                text = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', line)
                story.append(Paragraph(text, style_normal))

        # 5. Build PDF
        doc.build(story)
        return True
        
    except Exception as e:
        print(f"PDF Generation Error: {e}")
        import traceback
        traceback.print_exc()
        return False

def generate_revised_docx(clauses_data, output_path):
    """
    Generates a clean DOCX from the revised clauses.
    clauses_data: List of dicts containing 'revised_text' or 'content'
    """
    # --- helpers ---
    def _strip_tags(text: str) -> str:
        # minimal HTML tag stripper for table cells
        text = re.sub(r"<\s*br\s*/?\s*>", "\n", text, flags=re.I)
        text = re.sub(r"<.*?>", "", text)
        return html.unescape(text).strip()

    def _add_html_table(doc, table_html: str):
        rows = re.findall(r"<tr[^>]*>(.*?)</tr>", table_html, flags=re.S | re.I)
        parsed_rows = []
        max_cols = 0
        for row in rows:
            cells = re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", row, flags=re.S | re.I)
            cells = [_strip_tags(c) for c in cells]
            max_cols = max(max_cols, len(cells))
            parsed_rows.append(cells)
        if not parsed_rows or max_cols == 0:
            # fallback：无法解析则作为纯文本
            doc.add_paragraph(_strip_tags(table_html))
            return
        table = doc.add_table(rows=len(parsed_rows), cols=max_cols)
        table.style = "Table Grid"
        for i, row in enumerate(parsed_rows):
            for j, cell in enumerate(row):
                table.cell(i, j).text = cell

    def _add_markdown_table(doc, text: str):
        lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
        # 寻找表头和分隔线
        header_idx = None
        sep_idx = None
        for idx, ln in enumerate(lines):
            if "|" in ln:
                if header_idx is None:
                    header_idx = idx
                elif re.match(r"^\|?\s*:?-{3,}.*", ln):
                    sep_idx = idx
                    break
        if header_idx is None or sep_idx is None:
            doc.add_paragraph(text)
            return
        header = [c.strip() for c in lines[header_idx].strip("|").split("|")]
        body_lines = lines[sep_idx + 1 :]
        rows = []
        for ln in body_lines:
            if "|" not in ln:
                continue
            rows.append([c.strip() for c in ln.strip("|").split("|")])
        if not rows:
            doc.add_paragraph(text)
            return
        table = doc.add_table(rows=1 + len(rows), cols=len(header))
        table.style = "Table Grid"
        for j, cell in enumerate(header):
            table.cell(0, j).text = cell
        for i, row in enumerate(rows, start=1):
            for j, cell in enumerate(row):
                if j < len(header):
                    table.cell(i, j).text = cell

    def _add_clause_block(doc, text: str):
        # 统一预处理：解码 HTML 实体 & 反斜杠转义
        text = html.unescape(text or "")
        text = text.replace("\\<", "<").replace("\\>", ">")

        # 优先处理 HTML 表格
        if re.search(r"<table", text, flags=re.I):
            _add_html_table(doc, text)
            return
        # 其次处理 Markdown 表格
        if ("|" in text and "---" in text):
            _add_markdown_table(doc, text)
            return
        # 普通文本（保证不残留标签）
        doc.add_paragraph(_strip_tags(text))

    doc = Document()
    
    # Set default font to Microsoft YaHei
    style = doc.styles['Normal']
    style.font.name = 'Microsoft YaHei'
    from docx.oxml.ns import qn
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    
    doc.add_heading('修订后合同', 0)
    
    for clause in clauses_data:
        # Use revised_text if available, else content
        text = clause.get('revised_text', clause.get('content', ''))
        
        # Simple heuristic for headings based on 'type' or 'section' could be added here
        # For now, just add as paragraphs
        if clause.get('type') in ['首部', '尾部']:
             _add_clause_block(doc, text)
        else:
             # Add section title if it's a new section? 
             # For simplicity, we just dump the text. 
             # A better approach would be to reconstruct the document structure.
             # But since we have a flat list of clauses, we just add them.
             _add_clause_block(doc, text)
    
    doc.save(output_path)
    return True

def generate_redline_html(clauses_data, output_path):
    """
    Generates an HTML file showing the redline diffs.
    Uses the advanced layout from test_CR.py
    """
    contract_html_parts = []
    analysis_html_parts = []
    last_section = None

    for clause in clauses_data:
        clause_id = f"clause-{clause.get('clause_number', '0')}"
        analysis_id = f"analysis-{clause_id}"
        
        # Unescape HTML content to ensure tables and other elements render correctly
        # diff_match_patch or other processes might have escaped the HTML tags
        clause_diff = html.unescape(clause.get('diff_html', clause.get('content', '')))
        
        analysis_text = clause.get('analysis', '').replace('\n', '<br>')
        section_title = clause.get('section', '')
        
        # Evidence Handling
        evidence_html = ""
        evidence_list = clause.get('evidence', [])
        if evidence_list and isinstance(evidence_list, list):
            for ev in evidence_list:
                source = ev.get('source', 'Unknown Source')
                source_name = os.path.basename(source)
                source_name = os.path.splitext(source_name)[0]
                content = ev.get('content', '')
                if len(content) > 120: content = content[:120] + "..."
                
                evidence_html += f'''
                <div style="margin-bottom: 8px; padding-bottom: 8px; border-bottom: 1px dashed #eee;">
                    <div style="display: flex; align-items: center; gap: 6px; margin-bottom: 3px;">
                        <span style="font-size: 12px;">⚖️</span>
                        <span title="Source: {source}" style="font-weight: 600; color: #2980b9; font-size: 11px; border-bottom: 1px dotted #2980b9;">{source_name}</span>
                    </div>
                    <div style="font-size: 11px; color: #555; line-height: 1.4; background: #fff; padding: 4px; border-radius: 3px;">{content}</div>
                </div>
                '''
        else:
            raw_evidence = clause.get('related_laws', '未检索到具体法条。')
            if len(raw_evidence) > 150: raw_evidence = raw_evidence[:150] + "..."
            evidence_html = f'<div style="color: #95a5a6; font-style: italic;">{raw_evidence}</div>'

        if section_title and section_title != last_section:
            contract_html_parts.append(f"""
            <div class="section-header" style="font-weight: bold; font-size: 1.1em; margin-top: 30px; margin-bottom: 15px; color: #2c3e50; border-bottom: 2px solid #eee; padding-bottom: 8px;">
                {section_title}
            </div>
            """)
            last_section = section_title
        
        contract_html_parts.append(f"""
        <div id="{clause_id}" class="clause-text-block" 
             onmouseenter="highlightRight('{analysis_id}')"
             onmouseleave="unhighlightRight('{analysis_id}')"
             style="margin-bottom: 12px; padding: 12px; border-radius: 6px; transition: all 0.2s; display: flex; gap: 12px; border: 1px solid transparent;">
            <div class="clause-number" style="font-weight: bold; color: #7f8c8d; min-width: 35px; font-size: 0.9em; padding-top: 2px;">
                {clause.get('clause_number', '')}
            </div>
            <div class="clause-body" style="font-family: 'Georgia', serif; line-height: 1.6; flex: 1; color: #34495e;">
                {clause_diff}
            </div>
        </div>
        """)
        
        risk_level = clause.get('risk_level', 'Low Risk')
        risk_color = "#c0392b" if "高" in risk_level or "High" in risk_level else "#f39c12" if "中" in risk_level or "Medium" in risk_level else "#27ae60"
        risk_bg = "#ffebee" if "高" in risk_level or "High" in risk_level else "#fef9e7" if "中" in risk_level or "Medium" in risk_level else "#e8f8f5"
        
        analysis_html_parts.append(f"""
        <div id="{analysis_id}" class="analysis-card" 
             onmouseenter="highlightLeft('{clause_id}')"
             onmouseleave="unhighlightLeft('{clause_id}')"
             style="background: #fff; border-left: 4px solid {risk_color}; padding: 15px; margin-bottom: 20px; box-shadow: 0 2px 8px rgba(0,0,0,0.05); border-radius: 0 6px 6px 0; opacity: 0.85; transition: all 0.3s; position: relative;">
            
            <div style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 10px;">
                <h4 style="margin: 0; color: {risk_color}; font-size: 14px; font-weight: 700;">
                    ⚠️ 风险分析 ({clause.get('clause_number', '')})
                </h4>
                <span style="font-size: 10px; background: {risk_bg}; color: {risk_color}; padding: 2px 6px; border-radius: 4px;">{risk_level}</span>
            </div>
            
            <div style="font-size: 13px; color: #555; line-height: 1.5; margin-bottom: 12px;">
                {analysis_text}
            </div>

            <div style="background: #f8f9fa; padding: 10px; border-radius: 6px; border: 1px solid #e0e0e0; margin-top: 12px;">
                <div style="font-size: 10px; font-weight: 700; color: #95a5a6; margin-bottom: 8px; text-transform: uppercase; letter-spacing: 0.5px; display: flex; align-items: center; gap: 5px;">
                    <span>📚</span> 法律依据 / Legal Basis
                </div>
                <div style="font-size: 11px; color: #666;">
                    {evidence_html}
                </div>
            </div>
        </div>
        """)

    full_contract_html = "\n".join(contract_html_parts)
    full_analysis_html = "\n".join(analysis_html_parts)

    html_content = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="utf-8">
        <title>合同审查报告</title>
        <style>
            body {{ background-color: #f4f6f7; font-family: 'Microsoft YaHei', sans-serif; }}
            .container {{ max-width: 96%; padding: 20px; margin: 0 auto; }}
            .report-layout {{ display: flex; gap: 40px; align-items: flex-start; }}
            .contract-view {{ 
                flex: 2; 
                background: white; 
                padding: 50px; 
                border-radius: 8px; 
                box-shadow: 0 4px 20px rgba(0,0,0,0.06); 
                min-height: 80vh;
            }}
            .analysis-sidebar {{ 
                flex: 1; 
                position: sticky; 
                top: 20px; 
                max-height: 95vh; 
                overflow-y: auto; 
                padding-right: 10px;
                padding-left: 5px;
            }}
            .analysis-sidebar::-webkit-scrollbar {{ width: 6px; }}
            .analysis-sidebar::-webkit-scrollbar-thumb {{ background: #bdc3c7; border-radius: 3px; }}
            .analysis-card:hover, .analysis-card.active {{ 
                opacity: 1; 
                transform: translateX(-8px); 
                box-shadow: 0 8px 25px rgba(231, 76, 60, 0.15); 
                z-index: 10;
            }}
            .clause-text-block:hover, .clause-text-block.active {{
                background-color: #fff8e1;
                border-color: #ffe082;
                box-shadow: 0 2px 8px rgba(0,0,0,0.05);
            }}
            del {{ background-color: #ffcdd2; color: #b71c1c; text-decoration: line-through; padding: 0 2px; border-radius: 2px; }}
            ins {{ background-color: #c8e6c9; color: #1b5e20; text-decoration: none; border-bottom: 2px solid #2e7d32; padding: 0 2px; border-radius: 2px; }}
        </style>
        <script>
            function highlightRight(analysisId) {{
                const target = document.getElementById(analysisId);
                if (target) {{
                    target.classList.add('active');
                    target.scrollIntoView({{ behavior: 'smooth', block: 'nearest' }});
                }}
            }}
            function unhighlightRight(analysisId) {{
                const target = document.getElementById(analysisId);
                if (target) target.classList.remove('active');
            }}
            function highlightLeft(clauseId) {{
                const target = document.getElementById(clauseId);
                if (target) {{
                    target.classList.add('active');
                    target.scrollIntoView({{ behavior: 'smooth', block: 'center' }});
                }}
            }}
            function unhighlightLeft(clauseId) {{
                const target = document.getElementById(clauseId);
                if (target) target.classList.remove('active');
            }}
        </script>
    </head>
    <body>
        <div class="container">
            <h1 style="text-align: center; color: #2c3e50; margin-bottom: 40px;">合同审查详细报告</h1>
            <div class="report-layout">
                <div class="contract-view">
                    {full_contract_html}
                </div>
                <div class="analysis-sidebar">
                    {full_analysis_html}
                </div>
            </div>
        </div>
    </body>
    </html>
    """
    
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(html_content)
    return True

def create_export_zip(export_dir, zip_path, original_file_path=None):
    """
    Zips all files in export_dir into zip_path.
    Optionally includes the original file.
    """
    with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
        # Add generated files
        for root, dirs, files in os.walk(export_dir):
            for file in files:
                if file == os.path.basename(zip_path):
                    continue
                file_path = os.path.join(root, file)
                arcname = os.path.relpath(file_path, export_dir)
                zipf.write(file_path, arcname)
        
        # Add original file if provided
        if original_file_path and os.path.exists(original_file_path):
            # Create a folder inside zip for original
            arcname = os.path.join("原始文件备份", os.path.basename(original_file_path))
            zipf.write(original_file_path, arcname)
            
    return True
